"""Importación y exportación de datos vía Excel (openpyxl)."""

from __future__ import annotations

import io
import unicodedata
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import db
from .allocation import parse_number


def _norm(s: str) -> str:
    s = "".join(c for c in unicodedata.normalize("NFD", str(s or "")) if unicodedata.category(c) != "Mn")
    return s.strip().lower()


# Sinónimos de cabeceras aceptados en la importación de documentos.
ALIASES = {
    "tipo": ["tipo", "clase", "documento"],
    "numero": ["numero", "num", "nº", "n factura", "numero factura", "numero documento", "referencia", "ref"],
    "fecha": ["fecha", "fecha factura", "fecha documento", "date"],
    "tercero": ["tercero", "proveedor", "cliente", "acreedor", "nombre", "razon social"],
    "concepto": ["concepto", "descripcion", "detalle", "texto", "glosa"],
    "importe": ["importe", "base", "base imponible", "importe base", "cuantia", "coste", "gasto", "amount", "valor", "importe total", "total"],
    "iva_pct": ["iva", "% iva", "iva %", "tipo iva", "porcentaje iva"],
    "iva_importe": ["cuota iva", "importe iva", "iva importe"],
    "total": ["total factura", "total documento", "importe con iva", "total con iva"],
    "cuenta": ["cuenta", "cuenta contable", "cta", "cuenta contab", "cod cuenta", "codigo cuenta"],
    "centro": ["centro", "centro de coste", "centro coste", "cc", "centro de beneficio", "centro coste/beneficio", "cost center"],
    "nif": ["nif", "cif", "nif/cif", "dni"],
}

TIPOS_VALIDOS = {"factura", "albaran", "albaranes", "apunte", "apuntes", "factura recibida", "factura emitida"}


def _map_headers(headers: List[str]) -> Dict[str, int]:
    """Devuelve {campo_interno: indice_columna}."""
    mapping: Dict[str, int] = {}
    norm_headers = [_norm(h) for h in headers]
    for campo, alias in ALIASES.items():
        for a in alias:
            if a in norm_headers:
                mapping[campo] = norm_headers.index(a)
                break
    return mapping


def preview_workbook(file_bytes: bytes, max_rows: int = 8) -> Dict[str, Any]:
    """Lee cabeceras y primeras filas para mostrar un mapeo previo."""
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return {"headers": [], "mapping": {}, "sample": [], "n_rows": 0}
    headers = [str(h) if h is not None else "" for h in rows[0]]
    mapping = _map_headers(headers)
    sample = [list(r) for r in rows[1:1 + max_rows]]
    return {
        "headers": headers,
        "mapping": mapping,
        "sample": sample,
        "n_rows": max(0, len(rows) - 1),
    }


def _clave_doc(numero, importe) -> str:
    return f"{_norm(numero)}|{round(float(importe or 0), 2)}"


def import_documentos(conn, file_bytes: bytes, tipo_defecto: str = "factura",
                      crear_maestros: bool = True, aplicar_reglas: bool = True,
                      omitir_duplicados: bool = True, cerrados=None) -> Dict[str, Any]:
    """Importa documentos desde un Excel. Crea cuentas/centros/terceros si no existen.

    Si un centro tiene una regla de reparto por defecto, se aplica automáticamente
    al documento importado (aplicar_reglas=True).
    """
    from .allocation import Allocator, Project

    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    cerrados = cerrados or set()
    result = {"importados": 0, "errores": [], "cuentas_creadas": 0,
              "centros_creados": 0, "terceros_creados": 0, "repartidos_auto": 0,
              "duplicados": 0, "cerrados_omitidos": 0}
    if len(rows) < 2:
        result["errores"].append("El archivo no tiene filas de datos.")
        return result

    headers = [str(h) if h is not None else "" for h in rows[0]]
    m = _map_headers(headers)
    if "importe" not in m:
        result["errores"].append(
            "No se encuentra una columna de importe. Cabeceras detectadas: " + ", ".join(headers)
        )
        return result

    # cachés de maestros existentes
    cuentas = {_norm(c["codigo"]): c["id"] for c in db.list_cuentas(conn)}
    centros = {_norm(c["codigo"]): c for c in db.list_centros(conn)}
    terceros = {_norm(t["nombre"]): t["id"] for t in db.list_terceros(conn)}
    existentes = {_clave_doc(d["numero"], d["importe"]) for d in db.list_documentos(conn)
                  if (d["numero"] or "").strip()}

    # allocator para reglas por defecto
    projs = [Project(codigo=p["codigo"], nombre=p["nombre"],
                     drivers={k: v for k, v in (p.get("drivers") or {}).items()})
             for p in db.list_proyectos(conn, solo_activos=True)]
    allocator = Allocator(projs)
    cod2id = {p["codigo"]: p["id"] for p in db.list_proyectos(conn)}

    for i, row in enumerate(rows[1:], start=2):
        def cell(campo):
            idx = m.get(campo)
            return row[idx] if idx is not None and idx < len(row) else None

        imp_raw = cell("importe")
        importe = parse_number(imp_raw) if imp_raw is not None else None
        if importe is None:
            if all(c is None or str(c).strip() == "" for c in row):
                continue  # fila vacía
            result["errores"].append(f"Fila {i}: importe inválido ({imp_raw!r}).")
            continue

        numero_val = str(cell("numero") or "").strip()
        if omitir_duplicados and numero_val:
            clave = _clave_doc(numero_val, importe)
            if clave in existentes:
                result["duplicados"] += 1
                continue
            existentes.add(clave)

        # fecha normalizada y control de periodo cerrado
        fecha = cell("fecha")
        if fecha is not None and hasattr(fecha, "strftime"):
            fecha = fecha.strftime("%Y-%m-%d")
        fecha = str(fecha or "").strip()
        ej, per = db.periodo_desde_fecha(fecha)
        if cerrados and ej and per and (ej, per) in cerrados:
            result["cerrados_omitidos"] += 1
            continue

        tipo = _norm(cell("tipo")) or tipo_defecto
        if "albaran" in tipo:
            tipo = "albaran"
        elif "apunte" in tipo:
            tipo = "apunte"
        elif "factura" in tipo:
            tipo = "factura"
        else:
            tipo = tipo_defecto

        # cuenta
        cuenta_id = None
        cuenta_cod = cell("cuenta")
        if cuenta_cod is not None and str(cuenta_cod).strip():
            key = _norm(cuenta_cod)
            cuenta_id = cuentas.get(key)
            if cuenta_id is None and crear_maestros:
                cuenta_id = db.upsert_cuenta(conn, str(cuenta_cod).strip(), "")
                cuentas[key] = cuenta_id
                result["cuentas_creadas"] += 1

        # centro (la caché guarda el registro completo para leer su regla por defecto)
        centro_id = None
        centro = None
        centro_cod = cell("centro")
        if centro_cod is not None and str(centro_cod).strip():
            key = _norm(centro_cod)
            centro = centros.get(key)
            if centro is None and crear_maestros:
                cid = db.upsert_centro(conn, str(centro_cod).strip(), "")
                centro = db.get_centro(conn, cid)
                centros[key] = centro
                result["centros_creados"] += 1
            centro_id = centro["id"] if centro else None

        # tercero (maestro)
        tercero_txt = str(cell("tercero") or "").strip()
        tercero_id = None
        if tercero_txt:
            key = _norm(tercero_txt)
            tercero_id = terceros.get(key)
            if tercero_id is None and crear_maestros:
                nif = str(cell("nif") or "").strip()
                tercero_id = db.upsert_tercero(conn, tercero_txt, nif=nif)
                terceros[key] = tercero_id
                result["terceros_creados"] += 1

        # IVA
        iva_pct = parse_number(cell("iva_pct")) if cell("iva_pct") is not None else None
        iva_importe = parse_number(cell("iva_importe")) if cell("iva_importe") is not None else None

        did = db.insert_documento(
            conn,
            tipo=tipo,
            numero=numero_val,
            fecha=fecha,
            tercero=tercero_txt,
            tercero_id=tercero_id,
            concepto=str(cell("concepto") or "").strip(),
            importe=float(importe),
            iva_pct=float(iva_pct) if iva_pct is not None else 0,
            iva_importe=float(iva_importe) if iva_importe is not None else None,
            cuenta_id=cuenta_id,
            centro_id=centro_id,
            estado="pendiente",
        )
        result["importados"] += 1

        # reparto automático por regla por defecto del centro
        if aplicar_reglas and centro and (centro.get("regla_defecto") or "").strip() and projs:
            res = allocator.allocate(float(importe), centro["regla_defecto"])
            if res.ok and res.lines:
                lines = [{
                    "proyecto_id": cod2id.get(l.proyecto),
                    "proyecto_codigo": l.proyecto,
                    "importe": float(l.importe),
                    "porcentaje": float(l.porcentaje),
                    "base": f"[auto {centro['codigo']}] {centro['regla_defecto']} · {l.base}",
                } for l in res.lines]
                db.replace_repartos(conn, did, lines)
                result["repartidos_auto"] += 1

    return result


def import_maestro(conn, file_bytes: bytes, clase: str) -> Dict[str, Any]:
    """Importa proyectos / centros / cuentas desde un Excel simple (codigo, nombre...)."""
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    result = {"importados": 0, "errores": []}
    if len(rows) < 2:
        result["errores"].append("El archivo no tiene filas de datos.")
        return result
    headers = [_norm(h) for h in rows[0]]

    def idx(*names):
        for n in names:
            if n in headers:
                return headers.index(n)
        return None

    i_cod = idx("codigo", "cod", "code")
    i_nom = idx("nombre", "descripcion", "name")
    i_extra = idx("tipo") if clase == "centros" else idx("grupo", "presupuesto")
    if i_cod is None:
        result["errores"].append("Falta la columna 'codigo'.")
        return result

    for r in rows[1:]:
        cod = r[i_cod] if i_cod < len(r) else None
        if cod is None or str(cod).strip() == "":
            continue
        nombre = str(r[i_nom]).strip() if i_nom is not None and i_nom < len(r) and r[i_nom] else ""
        cod = str(cod).strip()
        if clase == "proyectos":
            presupuesto = 0
            ip = idx("presupuesto")
            if ip is not None and ip < len(r):
                presupuesto = float(parse_number(r[ip]) or 0)
            db.upsert_proyecto(conn, cod, nombre, presupuesto=presupuesto)
        elif clase == "centros":
            tipo = "coste"
            it = idx("tipo")
            if it is not None and it < len(r) and r[it]:
                tipo = "beneficio" if "benef" in _norm(r[it]) else "coste"
            db.upsert_centro(conn, cod, nombre, tipo=tipo)
        elif clase == "cuentas":
            grupo = ""
            ig = idx("grupo")
            if ig is not None and ig < len(r) and r[ig]:
                grupo = str(r[ig]).strip()
            db.upsert_cuenta(conn, cod, nombre, grupo=grupo)
        result["importados"] += 1
    return result


# --- Exportación ----------------------------------------------------------
# Cabeceras con la identidad corporativa: negro con texto verde lima.
_HEAD_FILL = PatternFill("solid", fgColor="101211")
_HEAD_FONT = Font(bold=True, color="8FDC1F")


def _style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = _HEAD_FILL
        cell.font = _HEAD_FONT
        cell.alignment = Alignment(horizontal="center")


def _autosize(ws):
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 10), 45)


def plantilla_documentos() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Documentos"
    headers = ["Tipo", "Numero", "Fecha", "Tercero", "NIF", "Concepto",
               "Importe", "IVA", "Cuenta", "Centro"]
    ws.append(headers)
    ejemplos = [
        ["factura", "F-2026/001", "2026-01-15", "Suministros Levante SL", "B96000001", "Material de obra", 3630.50, 21, "600000", "CC-OBRAS"],
        ["albaran", "ALB-1042", "2026-01-20", "Transportes Mediterraneo", "B96000002", "Portes enero", 480.00, 21, "624000", "CC-LOG"],
        ["apunte", "AS-5501", "2026-01-31", "Nomina personal tecnico", "", "Personal indirecto", 12500.00, 0, "640000", "CC-ESTRUCTURA"],
    ]
    for e in ejemplos:
        ws.append(e)
    _style_header(ws, len(headers))
    _autosize(ws)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def plantilla_maestro(clase: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    if clase == "proyectos":
        ws.title = "Proyectos"
        ws.append(["Codigo", "Nombre", "Presupuesto"])
        ws.append(["PROY-A", "Reforma Nave Norte", 150000])
        ws.append(["PROY-B", "Ampliacion Oficinas", 80000])
    elif clase == "centros":
        ws.title = "Centros"
        ws.append(["Codigo", "Nombre", "Tipo"])
        ws.append(["CC-OBRAS", "Obras y ejecucion", "coste"])
        ws.append(["CB-VENTAS", "Ventas", "beneficio"])
    else:
        ws.title = "Cuentas"
        ws.append(["Codigo", "Nombre", "Grupo"])
        ws.append(["600000", "Compras de mercaderias", "6"])
        ws.append(["640000", "Sueldos y salarios", "6"])
    _style_header(ws, ws.max_column)
    _autosize(ws)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def export_informe(conn) -> bytes:
    """Genera un Excel con el informe analítico completo."""
    wb = Workbook()

    # Hoja 1: resumen por proyecto
    ws = wb.active
    ws.title = "Por proyecto"
    ws.append(["Codigo", "Proyecto", "Presupuesto", "Imputado", "Desviacion", "% consumo", "Nº lineas"])
    for p in db.resumen_por_proyecto(conn):
        pres = p["presupuesto"] or 0
        imp = p["imputado"] or 0
        desv = pres - imp
        pct = (imp / pres * 100) if pres else 0
        ws.append([p["codigo"], p["nombre"], round(pres, 2), round(imp, 2),
                   round(desv, 2), round(pct, 1), p["n_lineas"]])
    _style_header(ws, 7)
    _autosize(ws)

    # Hoja 2: resumen por centro
    ws2 = wb.create_sheet("Por centro")
    ws2.append(["Codigo", "Centro", "Tipo", "Importe", "Nº documentos"])
    for c in db.resumen_por_centro(conn):
        ws2.append([c["codigo"], c["nombre"], c["tipo"], round(c["importe"] or 0, 2), c["n_docs"]])
    _style_header(ws2, 5)
    _autosize(ws2)

    # Hoja 3: detalle de repartos (líneas analíticas)
    ws3 = wb.create_sheet("Detalle repartos")
    ws3.append(["Proyecto", "Documento", "Fecha", "Tercero", "Concepto",
                "Tipo", "Centro", "Importe doc.", "Importe imputado", "%", "Criterio"])
    for p in db.list_proyectos(conn):
        for d in db.detalle_proyecto(conn, p["id"]):
            ws3.append([p["codigo"], d["numero"], d["fecha"], d["tercero"], d["concepto"],
                        d["tipo"], d["centro_codigo"], round(d["doc_importe"] or 0, 2),
                        round(d["importe"] or 0, 2), round(d["porcentaje"] or 0, 1), d["base"]])
    _style_header(ws3, 11)
    _autosize(ws3)

    # Hoja 4: resumen por cuenta contable
    ws4 = wb.create_sheet("Por cuenta")
    ws4.append(["Codigo", "Cuenta", "Grupo", "Importe", "Nº documentos"])
    for c in db.resumen_por_cuenta(conn):
        ws4.append([c["codigo"], c["nombre"], c["grupo"], round(c["importe"] or 0, 2), c["n_docs"]])
    _style_header(ws4, 5)
    _autosize(ws4)

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def export_documentos(conn, documentos: List[Dict[str, Any]]) -> bytes:
    """Exporta una lista de documentos (ya filtrada) a Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Documentos"
    ws.append(["Tipo", "Numero", "Fecha", "Tercero", "Concepto", "Importe",
               "IVA %", "Cuota IVA", "Total", "Cuenta", "Centro", "Estado"])
    for d in documentos:
        ws.append([d.get("tipo"), d.get("numero"), d.get("fecha"), d.get("tercero"),
                   d.get("concepto"), round(d.get("importe") or 0, 2),
                   round(d.get("iva_pct") or 0, 2), round(d.get("iva_importe") or 0, 2),
                   round(d.get("total") or 0, 2), d.get("cuenta_codigo"),
                   d.get("centro_codigo"), d.get("estado")])
    _style_header(ws, 12)
    _autosize(ws)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def export_backup(conn) -> bytes:
    """Copia de seguridad completa en JSON."""
    import json

    data = {
        "version": 2,
        "proyectos": db.list_proyectos(conn),
        "centros": db.list_centros(conn),
        "cuentas": db.list_cuentas(conn),
        "terceros": db.list_terceros(conn),
        "reglas": db.list_reglas(conn),
        "documentos": db.list_documentos(conn),
    }
    # repartos por documento
    reps = []
    for d in data["documentos"]:
        for r in db.get_repartos(conn, d["id"]):
            reps.append(r)
    data["repartos"] = reps
    return json.dumps(data, ensure_ascii=False, indent=2, default=str).encode("utf-8")


def restore_backup(conn, file_bytes: bytes) -> Dict[str, Any]:
    """Restaura una copia de seguridad JSON. SUSTITUYE todos los datos actuales."""
    import json

    try:
        data = json.loads(file_bytes.decode("utf-8"))
    except (ValueError, UnicodeDecodeError) as e:
        return {"ok": False, "error": f"Archivo JSON no válido: {e}"}
    if not isinstance(data, dict) or "documentos" not in data:
        return {"ok": False, "error": "El archivo no parece una copia de CostControl."}

    res = {"ok": True, "error": ""}
    try:
        with conn:
            for t in ("repartos", "documentos", "reglas", "terceros",
                      "cuentas", "centros", "proyectos"):
                conn.execute(f"DELETE FROM {t}")

            for p in data.get("proyectos", []):
                conn.execute(
                    "INSERT INTO proyectos (id,codigo,nombre,descripcion,activo,presupuesto,drivers) "
                    "VALUES (?,?,?,?,?,?,?)",
                    (p.get("id"), p.get("codigo"), p.get("nombre", ""), p.get("descripcion", ""),
                     int(p.get("activo", 1)), float(p.get("presupuesto", 0) or 0),
                     json.dumps(p.get("drivers") or {})))
            for c in data.get("centros", []):
                conn.execute(
                    "INSERT INTO centros (id,codigo,nombre,tipo,descripcion,regla_defecto) "
                    "VALUES (?,?,?,?,?,?)",
                    (c.get("id"), c.get("codigo"), c.get("nombre", ""), c.get("tipo", "coste"),
                     c.get("descripcion", ""), c.get("regla_defecto", "")))
            for c in data.get("cuentas", []):
                conn.execute("INSERT INTO cuentas (id,codigo,nombre,grupo) VALUES (?,?,?,?)",
                             (c.get("id"), c.get("codigo"), c.get("nombre", ""), c.get("grupo", "")))
            for t in data.get("terceros", []):
                conn.execute("INSERT INTO terceros (id,nombre,nif,tipo) VALUES (?,?,?,?)",
                             (t.get("id"), t.get("nombre", ""), t.get("nif", ""), t.get("tipo", "proveedor")))
            for r in data.get("reglas", []):
                conn.execute("INSERT INTO reglas (id,nombre,texto) VALUES (?,?,?)",
                             (r.get("id"), r.get("nombre", "Regla"), r.get("texto", "")))
            for d in data.get("documentos", []):
                conn.execute(
                    "INSERT INTO documentos (id,tipo,numero,fecha,tercero,tercero_id,concepto,"
                    "importe,iva_pct,iva_importe,total,ejercicio,periodo,cuenta_id,centro_id,"
                    "estado,notas,adjunto) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (d.get("id"), d.get("tipo", "factura"), d.get("numero", ""), d.get("fecha", ""),
                     d.get("tercero", ""), d.get("tercero_id"), d.get("concepto", ""),
                     float(d.get("importe", 0) or 0), float(d.get("iva_pct", 0) or 0),
                     float(d.get("iva_importe", 0) or 0), float(d.get("total", 0) or 0),
                     d.get("ejercicio"), d.get("periodo"), d.get("cuenta_id"), d.get("centro_id"),
                     d.get("estado", "pendiente"), d.get("notas", ""), d.get("adjunto", "")))
            for r in data.get("repartos", []):
                conn.execute(
                    "INSERT INTO repartos (id,documento_id,proyecto_id,proyecto_codigo,"
                    "importe,porcentaje,base) VALUES (?,?,?,?,?,?,?)",
                    (r.get("id"), r.get("documento_id"), r.get("proyecto_id"),
                     r.get("proyecto_codigo", ""), float(r.get("importe", 0) or 0),
                     float(r.get("porcentaje", 0) or 0), r.get("base", "")))
        res.update({
            "proyectos": len(data.get("proyectos", [])),
            "centros": len(data.get("centros", [])),
            "cuentas": len(data.get("cuentas", [])),
            "terceros": len(data.get("terceros", [])),
            "documentos": len(data.get("documentos", [])),
            "repartos": len(data.get("repartos", [])),
        })
    except Exception as e:  # noqa
        return {"ok": False, "error": f"Error al restaurar: {e}"}
    return res
